import io
import uuid
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from unittest.mock import AsyncMock, patch

from conftest import TestSessionLocal
from docx import Document as DocxDocument
from openpyxl import Workbook
from sqlalchemy import select

from app.models.document import DocumentChunk
from app.models.user import User, UserRole
from app.services.auth_service import hash_password


def _unique_slug(prefix: str = "tenant") -> str:
    """Genera un slug unico de organizacion para evitar colisiones entre tests."""
    return f"{prefix}-{uuid.uuid4().hex[:8]}"


async def _register(client, slug: str | None = None, email: str = "admin@example.com", password: str = "SecurePass123!"):
    slug = slug or _unique_slug()
    payload = {
        "tenant_name": "Empresa de Pruebas SAS",
        "tenant_slug": slug,
        "admin_email": email,
        "admin_password": password,
    }
    response = await client.post("/auth/register", json=payload)
    return response, slug


async def _login(client, email: str, password: str, tenant_slug: str | None = None):
    payload = {"email": email, "password": password}
    if tenant_slug:
        payload["tenant_slug"] = tenant_slug
    return await client.post("/auth/login", json=payload)


async def _register_and_login(client, email: str | None = None, password: str = "SecurePass123!"):
    """Registra un tenant + admin nuevo, inicia sesion y retorna los headers de Authorization y el slug."""
    email = email or f"user-{uuid.uuid4().hex[:8]}@example.com"
    _, slug = await _register(client, email=email, password=password)
    login_response = await _login(client, email, password, slug)
    token = login_response.json()["data"]["access_token"]
    return {"Authorization": f"Bearer {token}"}, slug


async def _get_tenant_id(client, headers: dict) -> uuid.UUID:
    response = await client.get("/auth/me", headers=headers)
    return uuid.UUID(response.json()["data"]["tenant_id"])


async def _create_additional_user(tenant_id: uuid.UUID, email: str, password: str, role: UserRole = UserRole.VIEWER) -> User:
    """Crea un segundo usuario directamente en DB para un tenant existente (no hay endpoint publico para esto)."""
    async with TestSessionLocal() as db:
        user = User(tenant_id=tenant_id, email=email, hashed_password=hash_password(password), role=role)
        db.add(user)
        await db.commit()
        await db.refresh(user)
        return user


def _make_minimal_pdf_bytes(text: str) -> bytes:
    """Construye un PDF minimo valido con un unico texto, sin depender de librerias de generacion de PDF."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 400 200] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream_content = f"BT /F1 12 Tf 50 150 Td ({text}) Tj ET".encode("latin-1")
    objects.append(
        b"<< /Length " + str(len(stream_content)).encode() + b" >>\nstream\n" + stream_content + b"\nendstream"
    )

    buf = bytearray()
    buf += b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(buf))
        buf += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref_offset = len(buf)
    buf += f"xref\n0 {len(objects) + 1}\n".encode()
    buf += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        buf += f"{off:010d} 00000 n \n".encode()
    buf += b"trailer\n"
    buf += f"<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    buf += b"startxref\n"
    buf += f"{xref_offset}\n".encode()
    buf += b"%%EOF"

    return bytes(buf)


def _make_docx_bytes(text: str) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


VALID_PDF_TEXT = "Este es un documento de prueba sobre normativa ambiental colombiana y su aplicacion."
VALID_PDF_BYTES = _make_minimal_pdf_bytes(VALID_PDF_TEXT)
VALID_DOCX_BYTES = _make_docx_bytes(
    "Este es un documento DOCX de prueba sobre normativa ambiental colombiana y su aplicacion practica."
)
VALID_TXT_BYTES = (
    b"Contenido de prueba con mas de cincuenta caracteres sobre normativa ambiental colombiana. " * 3
)


def _make_mhtml_bytes(html: str) -> bytes:
    """Construye un .mhtml minimo (HTML empaquetado en MIME multipart/related, igual que
    'Guardar como...' del navegador)."""
    message = MIMEMultipart("related")
    message.attach(MIMEText(html, "html", "utf-8"))
    return message.as_bytes()


def _make_xlsx_bytes(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


VALID_MHTML_BYTES = _make_mhtml_bytes(
    "<html><body><h1>Resolucion de prueba</h1>"
    "<p>Pagina guardada sobre normativa ambiental colombiana y su aplicacion practica.</p>"
    "</body></html>"
)
VALID_HTML_BYTES = (
    "<html><body><h1>Resolucion de prueba</h1>"
    "<p>Documento HTML de prueba sobre normativa ambiental colombiana y su aplicacion practica.</p>"
    "</body></html>"
).encode("utf-8")
VALID_XLSX_BYTES = _make_xlsx_bytes(
    [
        ["articulo", "descripcion"],
        ["1", "Normativa ambiental colombiana de prueba con suficientes caracteres para pasar el minimo"],
        ["2", "Segunda fila de prueba sobre normativa ambiental colombiana y su aplicacion practica"],
    ]
)
VALID_CSV_BYTES = (
    "articulo,descripcion\n"
    "1,Normativa ambiental colombiana de prueba con suficientes caracteres para pasar el minimo\n"
    "2,Segunda fila de prueba sobre normativa ambiental colombiana y su aplicacion practica\n"
).encode("utf-8")


def _patched_ingestion_embeddings(dim: int = 768, value: float = 0.1):
    """Mockea openrouter_client.create_embeddings con vectores deterministicos, sin llamadas reales."""

    async def _create_embeddings(texts, model, dimensions):
        return [[value] * dim for _ in texts]

    return patch(
        "app.services.openrouter_client.create_embeddings", new=AsyncMock(side_effect=_create_embeddings)
    )


class TestDocumentUpload:
    async def test_upload_valid_pdf_returns_202_processing(self, client):
        """1. Upload de PDF valido debe retornar 202 con status='processing'."""
        headers, _ = await _register_and_login(client)

        with _patched_ingestion_embeddings():
            response = await client.post(
                "/api/v1/documents/upload",
                files={"file": ("normativa.pdf", VALID_PDF_BYTES, "application/pdf")},
                data={"es_base_corpus": "false"},
                headers=headers,
            )

        body = response.json()
        assert response.status_code == 202
        assert body["success"] is True
        assert body["data"]["status"] == "processing"
        assert body["data"]["filename"] == "normativa.pdf"

    async def test_upload_invalid_extension_returns_422(self, client):
        """2. Upload con extension invalida (.exe) debe retornar 422."""
        headers, _ = await _register_and_login(client)

        response = await client.post(
            "/api/v1/documents/upload",
            files={"file": ("instalador.exe", b"contenido cualquiera", "application/octet-stream")},
            data={"es_base_corpus": "false"},
            headers=headers,
        )

        assert response.status_code == 422

    async def test_upload_file_too_large_returns_422(self, client):
        """3. Upload con archivo mayor a 50MB debe retornar 422."""
        headers, _ = await _register_and_login(client)

        oversized = b"a" * (50 * 1024 * 1024 + 1)
        response = await client.post(
            "/api/v1/documents/upload",
            files={"file": ("grande.txt", oversized, "text/plain")},
            data={"es_base_corpus": "false"},
            headers=headers,
        )

        assert response.status_code == 422

    async def test_list_documents_only_shows_own_tenant(self, client):
        """4. Listar documentos del tenant solo debe mostrar los propios."""
        headers_a, _ = await _register_and_login(client, email="doc-a@example.com")
        headers_b, _ = await _register_and_login(client, email="doc-b@example.com")

        with _patched_ingestion_embeddings():
            await client.post(
                "/api/v1/documents/upload",
                files={"file": ("doc_a.txt", VALID_TXT_BYTES, "text/plain")},
                data={"es_base_corpus": "false"},
                headers=headers_a,
            )

        response_b = await client.get("/api/v1/documents/", headers=headers_b)
        assert response_b.status_code == 200
        assert response_b.json()["data"]["items"] == []

        response_a = await client.get("/api/v1/documents/", headers=headers_a)
        assert response_a.json()["data"]["total"] == 1

    async def test_get_document_from_other_tenant_returns_404(self, client):
        """5. GET documento por ID de otro tenant debe retornar 404."""
        headers_a, _ = await _register_and_login(client, email="doc-c@example.com")
        headers_b, _ = await _register_and_login(client, email="doc-d@example.com")

        with _patched_ingestion_embeddings():
            upload_response = await client.post(
                "/api/v1/documents/upload",
                files={"file": ("doc_c.txt", VALID_TXT_BYTES, "text/plain")},
                data={"es_base_corpus": "false"},
                headers=headers_a,
            )
        document_id = upload_response.json()["data"]["id"]

        response = await client.get(f"/api/v1/documents/{document_id}", headers=headers_b)
        assert response.status_code == 404

    async def test_delete_document_removes_document_and_chunks(self, client):
        """6. DELETE documento debe eliminar el documento y sus chunks (chunk_count=0 en DB tras el borrado)."""
        headers, _ = await _register_and_login(client, email="doc-e@example.com")

        with _patched_ingestion_embeddings():
            upload_response = await client.post(
                "/api/v1/documents/upload",
                files={"file": ("doc_e.txt", VALID_TXT_BYTES, "text/plain")},
                data={"es_base_corpus": "false"},
                headers=headers,
            )
        document_id = uuid.UUID(upload_response.json()["data"]["id"])

        detail_response = await client.get(f"/api/v1/documents/{document_id}", headers=headers)
        assert detail_response.json()["data"]["status"] == "ready"
        assert detail_response.json()["data"]["chunk_count"] > 0

        delete_response = await client.delete(f"/api/v1/documents/{document_id}", headers=headers)
        assert delete_response.status_code == 200

        get_after_delete = await client.get(f"/api/v1/documents/{document_id}", headers=headers)
        assert get_after_delete.status_code == 404

        async with TestSessionLocal() as db:
            result = await db.execute(select(DocumentChunk).where(DocumentChunk.document_id == document_id))
            remaining_chunks = result.scalars().all()
        assert len(remaining_chunks) == 0


class TestIngestionService:
    async def test_upload_pdf_returns_202_with_processing_status(self, client):
        """7. Upload de PDF valido debe retornar 202 con status='processing', sin esperar el
        resultado final del pipeline async (que ahora clasifica y llama a OpenRouter, no a
        parse_document sync)."""
        headers, _ = await _register_and_login(client, email="ingest-pdf@example.com")

        with (
            patch(
                "app.services.ingestion_service.openrouter_client.chat_completion",
                new=AsyncMock(return_value="ARTICULO 1. " + VALID_PDF_TEXT),
            ),
            _patched_ingestion_embeddings(),
        ):
            response = await client.post(
                "/api/v1/documents/upload",
                files={"file": ("normativa.pdf", VALID_PDF_BYTES, "application/pdf")},
                data={"es_base_corpus": "false"},
                headers=headers,
            )

        body = response.json()
        assert response.status_code == 202
        assert body["data"]["status"] == "processing"
        assert body["data"]["filename"] == "normativa.pdf"

    def test_parse_document_docx_returns_non_empty_text(self):
        """8. parse_document con DOCX valido debe retornar un string no vacio."""
        from app.services.ingestion_service import parse_document

        text = parse_document(VALID_DOCX_BYTES, "docx")
        assert len(text.strip()) > 0
        assert "normativa" in text.lower()

    def test_parse_document_mhtml_extracts_text_without_html_tags(self):
        """8b. parse_document con MHTML debe extraer el texto de la parte text/html, sin tags."""
        from app.services.ingestion_service import parse_document

        text = parse_document(VALID_MHTML_BYTES, "mhtml")
        assert "normativa ambiental colombiana" in text.lower()
        assert "<p>" not in text and "<html>" not in text

    def test_parse_document_html_extracts_text_without_html_tags(self):
        """8c. parse_document con HTML plano debe extraer el texto, sin tags."""
        from app.services.ingestion_service import parse_document

        text = parse_document(VALID_HTML_BYTES, "html")
        assert "normativa ambiental colombiana" in text.lower()
        assert "<h1>" not in text

    def test_parse_document_xlsx_includes_sheet_name_and_cells(self):
        """8d. parse_document con XLSX debe incluir el nombre de la hoja y el contenido de las celdas."""
        from app.services.ingestion_service import parse_document

        text = parse_document(VALID_XLSX_BYTES, "xlsx")
        assert "Hoja:" in text
        assert "normativa ambiental colombiana" in text.lower()

    def test_parse_document_csv_joins_cells_with_pipe(self):
        """8e. parse_document con CSV debe unir las celdas de cada fila."""
        from app.services.ingestion_service import parse_document

        text = parse_document(VALID_CSV_BYTES, "csv")
        assert "articulo | descripcion" in text
        assert "normativa ambiental colombiana" in text.lower()

    async def test_upload_valid_xlsx_is_processed_end_to_end(self, client):
        """8f. Upload de XLSX valido debe terminar en status='ready' con chunks generados
        (pipeline completo, no solo el parser unitario)."""
        headers, _ = await _register_and_login(client, email="upload-xlsx@example.com")

        with _patched_ingestion_embeddings():
            upload_response = await client.post(
                "/api/v1/documents/upload",
                files={
                    "file": (
                        "hoja.xlsx",
                        VALID_XLSX_BYTES,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
                data={"es_base_corpus": "false"},
                headers=headers,
            )
        document_id = upload_response.json()["data"]["id"]

        detail_response = await client.get(f"/api/v1/documents/{document_id}", headers=headers)
        assert detail_response.json()["data"]["status"] == "ready"
        assert detail_response.json()["data"]["chunk_count"] > 0

    def test_chunk_semantic_splits_by_articulo_with_metadata(self):
        """9. chunk_semantic con texto de al menos 2 ARTICULOs debe generar al menos 2 chunks,
        cada uno con metadata de numero_articulo/pagina_origen/posicion_en_documento y ninguno
        por debajo del minimo de 100 caracteres configurado."""
        from app.services.ingestion_service import chunk_semantic

        filler = (
            "Este parrafo desarrolla el contenido del articulo con suficientes caracteres "
            "para superar el minimo configurado. "
        )
        text = "ARTICULO 1. " + filler * 3 + "\n\n" + "ARTICULO 2. " + filler * 3

        chunks = chunk_semantic([(1, text)])

        assert len(chunks) >= 2
        for chunk in chunks:
            assert {"numero_articulo", "pagina_origen", "posicion_en_documento"} <= chunk["metadata"].keys()
            assert len(chunk["content"]) >= 100


class TestRagRetrieval:
    async def test_retrieve_relevant_chunks_returns_chunks_of_correct_tenant(self, client):
        """10. retrieve_relevant_chunks debe retornar chunks del tenant que subio el documento."""
        headers, _ = await _register_and_login(client, email="rag-a@example.com")
        tenant_id = await _get_tenant_id(client, headers)

        with _patched_ingestion_embeddings():
            await client.post(
                "/api/v1/documents/upload",
                files={"file": ("doc_rag.txt", VALID_TXT_BYTES, "text/plain")},
                data={"es_base_corpus": "false"},
                headers=headers,
            )

        from app.services.rag_service import retrieve_relevant_chunks

        async with TestSessionLocal() as db:
            chunks = await retrieve_relevant_chunks(db, "normativa ambiental", tenant_id, top_k=5)

        assert len(chunks) > 0
        assert all(chunk["document_id"] is not None for chunk in chunks)


class TestConversations:
    async def test_create_conversation_returns_201_with_id(self, client):
        """11. Crear conversacion debe retornar 201 con id."""
        headers, _ = await _register_and_login(client, email="conv-a@example.com")

        response = await client.post(
            "/api/v1/conversations/", json={"title": "Consulta sobre licencias"}, headers=headers
        )
        body = response.json()

        assert response.status_code == 201
        assert body["success"] is True
        assert "id" in body["data"]

    async def test_send_message_with_mocked_gemini_returns_assistant_response(self, client):
        """12. Enviar mensaje en conversacion con mocks de OpenRouter debe retornar 200 con
        respuesta del assistant. Sube un documento antes para que haya chunks reales que
        recuperar (si no hay chunks, el flujo responde el fallback en vez de llamar al LLM)."""
        headers, _ = await _register_and_login(client, email="conv-b@example.com")

        with _patched_ingestion_embeddings():
            await client.post(
                "/api/v1/documents/upload",
                files={"file": ("doc_conv_b.txt", VALID_TXT_BYTES, "text/plain")},
                data={"es_base_corpus": "false"},
                headers=headers,
            )

        conv_response = await client.post("/api/v1/conversations/", json={}, headers=headers)
        conversation_id = conv_response.json()["data"]["id"]

        with patch(
            "app.services.rag_service.openrouter_client.chat_completion",
            new=AsyncMock(return_value="La normativa vigente establece requisitos especificos."),
        ):
            response = await client.post(
                f"/api/v1/conversations/{conversation_id}/messages",
                json={"content": "Cual es la normativa sobre vertimientos?"},
                headers=headers,
            )

        body = response.json()
        assert response.status_code == 200
        assert body["data"]["role"] == "assistant"
        assert body["data"]["content"] == "La normativa vigente establece requisitos especificos."

    async def test_conversation_history_returns_chronological_order(self, client):
        """13. Historial de conversacion debe retornar los mensajes en orden cronologico."""
        headers, _ = await _register_and_login(client, email="conv-c@example.com")

        with _patched_ingestion_embeddings():
            await client.post(
                "/api/v1/documents/upload",
                files={"file": ("doc_conv_c.txt", VALID_TXT_BYTES, "text/plain")},
                data={"es_base_corpus": "false"},
                headers=headers,
            )

        conv_response = await client.post("/api/v1/conversations/", json={}, headers=headers)
        conversation_id = conv_response.json()["data"]["id"]

        for reply in ("Respuesta 1", "Respuesta 2"):
            with patch(
                "app.services.rag_service.openrouter_client.chat_completion",
                new=AsyncMock(return_value=reply),
            ):
                await client.post(
                    f"/api/v1/conversations/{conversation_id}/messages",
                    json={"content": f"Pregunta para {reply}"},
                    headers=headers,
                )

        history_response = await client.get(f"/api/v1/conversations/{conversation_id}/messages", headers=headers)
        messages = history_response.json()["data"]

        assert len(messages) == 4
        timestamps = [m["created_at"] for m in messages]
        assert timestamps == sorted(timestamps)

    async def test_send_message_to_other_users_conversation_returns_403(self, client):
        """14. Enviar mensaje en conversacion de otro usuario del mismo tenant debe retornar 403."""
        headers_a, slug = await _register_and_login(client, email="conv-owner@example.com")
        tenant_id = await _get_tenant_id(client, headers_a)

        await _create_additional_user(tenant_id, "conv-intruder@example.com", "SecurePass123!")
        login_b = await _login(client, "conv-intruder@example.com", "SecurePass123!", slug)
        token_b = login_b.json()["data"]["access_token"]
        headers_b = {"Authorization": f"Bearer {token_b}"}

        conv_response = await client.post("/api/v1/conversations/", json={}, headers=headers_a)
        conversation_id = conv_response.json()["data"]["id"]

        response = await client.post(
            f"/api/v1/conversations/{conversation_id}/messages",
            json={"content": "Intento no autorizado"},
            headers=headers_b,
        )

        assert response.status_code == 403
