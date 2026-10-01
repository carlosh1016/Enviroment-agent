import asyncio
import base64
import io
import re
import uuid

import pdfplumber
from docx import Document as DocxDocument
from pdf2image import convert_from_bytes
from sqlalchemy.ext.asyncio import AsyncSession

from app import database as app_database
from app.logging_config import get_logger
from app.models.document import Document, DocumentChunk, DocumentStatus
from app.services import openrouter_client

logger = get_logger(__name__)

# qwen3-embedding esta entrenado con Matryoshka (MRL): recortar a 768 dims conserva la calidad
# y evita migrar la columna Vector(768). Cambiar de modelo obliga a re-ingestar todo el corpus.
EMBEDDING_MODEL = "qwen/qwen3-embedding-8b"
EMBEDDING_DIMENSIONS = 768
EMBEDDING_BATCH_SIZE = 32
EMBEDDING_MAX_CONCURRENT_REQUESTS = 5

# Semaforo global del proceso: varios documentos se procesan en paralelo (uno por background
# task) y todos comparten la cuenta de OpenRouter. openrouter_client ya reintenta los 429.
_embed_semaphore = asyncio.Semaphore(EMBEDDING_MAX_CONCURRENT_REQUESTS)

# PDFs con texto van con el texto crudo de cada pagina; PDFs escaneados (sin capa de texto)
# van pagina por pagina como imagen. Mismo modelo multimodal para ambos casos.
# Nota: "google/gemini-flash-1.5-8b" (pedido originalmente) ya no existe en el catalogo de
# OpenRouter -> 404 "No endpoints found" (confirmado en vivo). Gemini 1.5 esta descontinuado.
# Se uso el equivalente actual: un modelo Flash chico/barato con soporte multimodal.
OPENROUTER_MODEL = "google/gemini-3.5-flash-lite"
OPENROUTER_MAX_TOKENS = 4096
PDF_TEXT_MIN_CHARS = 50

PDF_TEXT_SYSTEM_PROMPT = (
    "Eres un asistente especializado en normativa ambiental colombiana. "
    "Recibiras texto extraido de un PDF legal. Tu tarea es: "
    "1. Eliminar artefactos de extraccion (saltos de linea extranos, caracteres basura, "
    "headers/footers repetidos) "
    "2. Preservar la estructura legal: numeros de articulo, parrafos, numerales, literales "
    "3. Normalizar la numeracion (ARTICULO 1, ARTICULO 2, etc.) "
    "4. Devolver SOLO el texto limpio y estructurado, sin comentarios"
)

PDF_IMAGE_SYSTEM_PROMPT = (
    "Eres un asistente especializado en normativa ambiental colombiana. "
    "Recibiras la imagen de una pagina de un documento legal escaneado. Tu tarea es: "
    "1. Transcribir TODO el texto visible con precision "
    "2. Preservar la estructura: articulos, parrafos, numerales, literales "
    "3. Normalizar la numeracion (ARTICULO 1, ARTICULO 2, etc.) "
    "4. Devolver SOLO el texto transcrito y estructurado, sin comentarios"
)

# Chunking semantico por estructura legal (ver chunk_semantic). Formato esperado tras el paso
# de limpieza/transcripcion via LLM: "ARTICULO <numero>" (normalizado por el prompt de arriba).
ARTICLE_PATTERN = re.compile(r"ART[IÍ]CULO\s+(\d+)", re.IGNORECASE)
CHUNK_MIN_CHARS = 100
CHUNK_MAX_CHARS = 1500


def parse_document(file_bytes: bytes, file_type: str) -> str:
    """Extrae el texto plano de un archivo DOCX o TXT. Lanza ValueError si queda vacio o ilegible.

    Los PDF NO pasan por aqui: usan el pipeline async de clasificacion + OpenRouter
    (ver extract_pdf_pages), porque requieren I/O de red para limpiar/transcribir el texto.
    """
    if file_type == "docx":
        text = _parse_docx(file_bytes)
    elif file_type == "txt":
        text = file_bytes.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Tipo de archivo no soportado en parse_document: {file_type}")

    if len(text.strip()) < 50:
        raise ValueError("Documento vacio o ilegible")

    return text


def _parse_docx(file_bytes: bytes) -> str:
    """Extrae texto de un DOCX, incluyendo parrafos y celdas de tablas."""
    doc = DocxDocument(io.BytesIO(file_bytes))
    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text]

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return "\n\n".join(parts)


def _extract_raw_pdf_pages(file_bytes: bytes) -> list[str]:
    """Extrae el texto crudo de cada pagina de un PDF con pdfplumber (sin limpiar)."""
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        return [page.extract_text() or "" for page in pdf.pages]


def _is_scanned_pdf(first_page_text: str) -> bool:
    """Clasificador (Paso 1): un PDF se considera escaneado si pdfplumber no logra extraer
    al menos PDF_TEXT_MIN_CHARS caracteres de su primera pagina."""
    return len(first_page_text.strip()) < PDF_TEXT_MIN_CHARS


async def _clean_text_page(raw_text: str) -> str:
    """Paso 2 (PDF con texto): limpia y estructura el texto crudo de una pagina via OpenRouter."""
    messages = [
        {"role": "system", "content": PDF_TEXT_SYSTEM_PROMPT},
        {"role": "user", "content": raw_text},
    ]
    return await openrouter_client.chat_completion(
        messages, model=OPENROUTER_MODEL, max_tokens=OPENROUTER_MAX_TOKENS
    )


async def _transcribe_image_page(image_bytes: bytes) -> str:
    """Paso 2 (PDF escaneado): transcribe la imagen de una pagina via el modelo multimodal de OpenRouter."""
    image_b64 = base64.b64encode(image_bytes).decode("ascii")
    messages = [
        {"role": "system", "content": PDF_IMAGE_SYSTEM_PROMPT},
        {
            "role": "user",
            "content": [
                {"type": "text", "text": "Transcribe el texto de esta pagina."},
                {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{image_b64}"}},
            ],
        },
    ]
    return await openrouter_client.chat_completion(
        messages, model=OPENROUTER_MODEL, max_tokens=OPENROUTER_MAX_TOKENS
    )


def _render_pdf_pages_to_png(file_bytes: bytes) -> list[bytes]:
    """Convierte cada pagina de un PDF escaneado a una imagen PNG (requiere poppler-utils)."""
    images = convert_from_bytes(file_bytes)
    rendered = []
    for image in images:
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        rendered.append(buffer.getvalue())
    return rendered


async def extract_pdf_pages(file_bytes: bytes) -> list[tuple[int, str]]:
    """Clasifica el PDF (Paso 1) y devuelve el texto limpio/transcrito de cada pagina no vacia
    via OpenRouter (Paso 2), como lista de (numero_de_pagina, texto), 1-indexado."""
    raw_pages = await asyncio.to_thread(_extract_raw_pdf_pages, file_bytes)
    if not raw_pages:
        return []

    if _is_scanned_pdf(raw_pages[0]):
        logger.info("PDF clasificado como escaneado, transcribiendo pagina por pagina via OpenRouter")
        page_images = await asyncio.to_thread(_render_pdf_pages_to_png, file_bytes)
        pages: list[tuple[int, str]] = []
        for page_num, image_bytes in enumerate(page_images, start=1):
            transcribed = await _transcribe_image_page(image_bytes)
            pages.append((page_num, transcribed))
        return pages

    logger.info("PDF clasificado como texto, limpiando pagina por pagina via OpenRouter")
    pages = []
    for page_num, raw_text in enumerate(raw_pages, start=1):
        if not raw_text.strip():
            continue
        cleaned = await _clean_text_page(raw_text)
        pages.append((page_num, cleaned))
    return pages


def _split_by_article_or_paragraph(text: str) -> list[tuple[str, str | None]]:
    """Separa un texto en segmentos (segmento, numero_articulo_o_None).

    Si el texto contiene marcas "ARTICULO N", separa por articulo. Si no, separa por
    parrafo (doble salto de linea).
    """
    matches = list(ARTICLE_PATTERN.finditer(text))
    if not matches:
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        return [(p, None) for p in paragraphs]

    segments: list[tuple[str, str | None]] = []

    preamble = text[: matches[0].start()].strip()
    if preamble:
        segments.append((preamble, None))

    for i, match in enumerate(matches):
        start = match.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        segment_text = text[start:end].strip()
        if segment_text:
            segments.append((segment_text, match.group(1)))

    return segments


def _enforce_chunk_size(segment_text: str, numero_articulo: str | None) -> list[tuple[str, str | None]]:
    """Aplica los limites de tamano a un segmento: descarta si es muy chico, lo deja igual si
    entra en el maximo, o lo parte por parrafo (empaquetando greedily) si lo excede."""
    if len(segment_text) < CHUNK_MIN_CHARS:
        return []
    if len(segment_text) <= CHUNK_MAX_CHARS:
        return [(segment_text, numero_articulo)]

    paragraphs = [p.strip() for p in segment_text.split("\n\n") if p.strip()]
    if len(paragraphs) <= 1:
        # Sin parrafos para partir: se corta duro cada CHUNK_MAX_CHARS caracteres.
        paragraphs = [segment_text[i : i + CHUNK_MAX_CHARS] for i in range(0, len(segment_text), CHUNK_MAX_CHARS)]

    result: list[tuple[str, str | None]] = []
    buffer = ""
    for paragraph in paragraphs:
        candidate = f"{buffer}\n\n{paragraph}" if buffer else paragraph
        if len(candidate) <= CHUNK_MAX_CHARS:
            buffer = candidate
        else:
            if len(buffer) >= CHUNK_MIN_CHARS:
                result.append((buffer, numero_articulo))
            buffer = paragraph

    if len(buffer) >= CHUNK_MIN_CHARS:
        result.append((buffer, numero_articulo))

    return result


def chunk_semantic(pages: list[tuple[int | None, str]]) -> list[dict]:
    """Chunking semantico por estructura legal (Paso 3): separa por ARTICULO cuando el texto
    tiene esa estructura, o por parrafo si no. Cada chunk queda entre CHUNK_MIN_CHARS y
    CHUNK_MAX_CHARS caracteres, con metadata de numero_articulo, pagina_origen y
    posicion_en_documento (indice global 1-based dentro del documento).
    """
    chunks: list[dict] = []
    position = 0

    for page_num, page_text in pages:
        if not page_text or not page_text.strip():
            continue

        for segment_text, numero_articulo in _split_by_article_or_paragraph(page_text):
            for final_text, articulo in _enforce_chunk_size(segment_text, numero_articulo):
                position += 1
                chunks.append(
                    {
                        "content": final_text,
                        "metadata": {
                            "numero_articulo": articulo,
                            "pagina_origen": page_num,
                            "posicion_en_documento": position,
                        },
                    }
                )

    return chunks


async def _embed_batch(batch: list[str]) -> list[list[float]]:
    """Genera embeddings de un batch bajo el semaforo global (los reintentos ante 429 o 5xx
    los maneja openrouter_client)."""
    async with _embed_semaphore:
        return await openrouter_client.create_embeddings(batch, EMBEDDING_MODEL, EMBEDDING_DIMENSIONS)


async def embed_chunks(chunks: list[str]) -> list[list[float]]:
    """Genera embeddings para una lista de chunks, en batches de EMBEDDING_BATCH_SIZE.

    Un semaforo global (_embed_semaphore) limita a EMBEDDING_MAX_CONCURRENT_REQUESTS batches
    concurrentes entre TODOS los documentos que se procesan en paralelo.
    """
    vectors: list[list[float]] = []
    for i in range(0, len(chunks), EMBEDDING_BATCH_SIZE):
        vectors.extend(await _embed_batch(chunks[i : i + EMBEDDING_BATCH_SIZE]))
    return vectors


async def process_document(
    db: AsyncSession,
    document_id: uuid.UUID,
    file_bytes: bytes,
    file_type: str,
    tenant_id: uuid.UUID,
) -> None:
    """Orquesta el pipeline de ingesta: extraccion (PDF via OpenRouter, DOCX/TXT local),
    chunking semantico, embeddings e insercion de chunks en bulk.

    Si cualquier paso falla, deja el documento en status='error' con el detalle en error_message.
    """
    document = await db.get(Document, document_id)
    if document is None:
        logger.error("Documento no encontrado al iniciar el procesamiento")
        return

    document.status = DocumentStatus.PROCESSING
    await db.commit()

    try:
        if file_type == "pdf":
            pages = await extract_pdf_pages(file_bytes)
        else:
            text = await asyncio.to_thread(parse_document, file_bytes, file_type)
            pages = [(None, text)]

        chunks = chunk_semantic(pages)
        if not chunks:
            raise ValueError("No se pudieron generar fragmentos del documento")

        vectors = await embed_chunks([chunk["content"] for chunk in chunks])

        for index, (chunk, vector) in enumerate(zip(chunks, vectors)):
            db.add(
                DocumentChunk(
                    tenant_id=tenant_id,
                    document_id=document_id,
                    chunk_index=index,
                    content=chunk["content"],
                    embedding=vector,
                    metadata_=chunk["metadata"],
                )
            )

        document.status = DocumentStatus.READY
        document.chunk_count = len(chunks)
        document.error_message = None
        await db.commit()

        logger.info("Documento procesado exitosamente")
    except Exception as exc:
        await db.rollback()
        document = await db.get(Document, document_id)
        if document is not None:
            document.status = DocumentStatus.ERROR
            document.error_message = str(exc)[:1000]
            await db.commit()
        logger.error("Error procesando documento: %s", type(exc).__name__)


async def process_document_background(
    document_id: uuid.UUID,
    file_bytes: bytes,
    file_type: str,
    tenant_id: uuid.UUID,
) -> None:
    """Wrapper para ejecutar process_document como BackgroundTask, con su propia sesion de base de datos.

    Referencia app_database.AsyncSessionLocal por atributo de modulo (no por import directo) para que
    los tests puedan sustituirlo por su propio session factory de SQLite en memoria.
    """
    async with app_database.AsyncSessionLocal() as db:
        await process_document(db, document_id, file_bytes, file_type, tenant_id)
